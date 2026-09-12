from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
root=Path(__file__).resolve().parent / 'tarkistuspaketti'
d=Document(); sec=d.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=d.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(10.5 if name=='Normal' else {'Title':22,'Subtitle':12,'Heading 1':17,'Heading 2':13,'Heading 3':11}[name])
 s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.08

for style in d.styles:
 for e in list(style.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
normal=d.styles['Normal'];normal.paragraph_format.widow_control=True
h=sec.header.paragraphs[0];h.text='HIETAKULMA  |  LUONNOS 0.1  |  SISÄINEN TARKISTUS';h.runs[0].font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('12.9.2026  ·  Sivu ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for r in f.runs:r.font.size=Pt(8)
def clean(s):return re.sub(r'\*\*|`','',s)
def table(rows):
 rows=[r for r in rows if not re.fullmatch(r'[\s|:\-]+',r)]
 data=[[clean(v.strip()) for v in row.strip().strip('|').split('|')] for row in rows]
 n=len(data[0]);t=d.add_table(rows=0, cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 weights={2:[.36,.64],3:[.27,.39,.34],4:[.23,.25,.27,.25],5:[.21,.26,.14,.24,.15],6:[.15,.16,.24,.14,.18,.13]}.get(n,[1/n]*n)
 if data[0][0]=='Tunnus':weights=[.10,.40,.50]
 if data[0][0]=='Elementti tai toimitusrivi':weights=[.18,.10,.27,.23,.22]
 if data[0][0]=='Aiempi nimike':weights=[.29,.53,.18]
 for col,w in zip(t.columns,weights):col.width=Cm(17.4*w)
 for i,values in enumerate(data):
  row=t.add_row()
  pr=row._tr.get_or_add_trPr();avoid=OxmlElement('w:cantSplit');pr.append(avoid)
  if i==0:
   repeat=OxmlElement('w:tblHeader');pr.append(repeat)
  for j,(cell,value) in enumerate(zip(row.cells,values)):
   cell.width=Cm(17.4*weights[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tc.append(b);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65' if path.name.startswith('02-') else '90');e.set(qn('w:type'),'dxa');m.append(e)
   tc.append(m)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6EDF2');tc.append(sh)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
   r=p.add_run(value);r.font.size=Pt(9);r.bold=i==0
 d.add_paragraph().paragraph_format.space_after=Pt(1)
files=[root/'README.md',root/'LAHTEET.md']+sorted(root.glob('[0-9]*.md'))
for index,path in enumerate(files):
 if index:d.add_page_break()
 lines=path.read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i];i+=1
  if not line.strip():continue
  if line.startswith('|'):
   rows=[line]
   while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
   table(rows);continue
  if line.startswith('# '):
   title=clean(line[2:]);d.add_paragraph(title,'Title' if index==0 else 'Heading 1');continue
  if line.startswith('## '):
   if path.name.startswith('07-') and line[3:] in ['Jorma Salomäelle','Tapani Katajistolle']:d.add_page_break()
   d.add_paragraph(clean(line[3:]),'Heading 2');continue
  p=d.add_paragraph(clean(line))
  if path.name.startswith('02-'):p.paragraph_format.space_after=Pt(5)
d.core_properties.title='Hietakulman dokumenttien ja liitteiden tarkistuspaketti'
d.core_properties.subject='Luonnos 0.1 Villen Jorman ja Tapanin tarkistettavaksi'
d.core_properties.author='Hietakulma'
out=root/'Hietakulma-tarkistuspaketti-luonnos.docx';d.save(out);print(out)
