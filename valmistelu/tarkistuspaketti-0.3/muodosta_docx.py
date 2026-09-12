from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
root=Path(__file__).resolve().parent
d=Document(); sec=d.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.7);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=d.styles[name];s.font.name='Work Sans';s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(10.5 if name=='Normal' else {'Title':22,'Subtitle':12,'Heading 1':17,'Heading 2':13,'Heading 3':11}[name])
 s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.08

for style in d.styles:
 for e in list(style.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
normal=d.styles['Normal'];normal.paragraph_format.widow_control=True
sec.header_distance=Cm(.65)
h=sec.header.paragraphs[0];h.add_run().add_picture(str(root/'assets/hietakulma-logo.png'),width=Cm(6.8));h.add_run('    LUONNOS 0.3  /  TARKISTETTAVAKSI').font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('13.9.2026  ·  Sivu ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for r in f.runs:r.font.size=Pt(8)
def clean(s):
 s=re.sub(r'\*\*|`','',s)
 for a,b in [('Tuulensuojakipsilevy','Tuulensuoja\u00adkipsilevy'),('Sähkövaraus','Sähkö\u00advaraus'),('Rakennetyyppi','Rakenne\u00adtyyppi'),('Tuotantopiirustus','Tuotanto\u00adpiirustus'),('Räystäselementti','Räystäs\u00adelementti'),('Terassielementti','Terassi\u00adelementti'),('suoritustasoilmoituksen','suoritustaso\u00adilmoituksen'),('Suoritustasoilmoituksen','Suoritustaso\u00adilmoituksen'),('Varmennustodistus','Varmennus\u00adtodistus'),('rakennustuotedirektiivi','rakennustuote\u00addirektiivi')]:s=s.replace(a,b)
 return s
def table(rows):
 rows=[r for r in rows if not re.fullmatch(r'[\s|:\-]+',r)]
 data=[[clean(v.strip()) for v in row.strip().strip('|').split('|')] for row in rows]
 n=len(data[0]);t=d.add_table(rows=0, cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 weights={2:[.36,.64],3:[.27,.39,.34],4:[.23,.25,.27,.25],5:[.21,.26,.14,.24,.15],6:[.15,.16,.24,.14,.18,.13]}.get(n,[1/n]*n)
 if data[0][0]=='Tunnus' and n==3:weights=[.10,.40,.50]
 if data[0][0]=='Liite' and n==5:weights=[.07,.25,.15,.24,.29]
 if data[0][0]=='Toimituksen osa' and n==4:weights=[.29,.36,.13,.22]
 if data[0][0]=='Osa' and n==3:weights=[.07,.51,.42] if data[0][1]=='Sisältö' else [.20,.47,.33]
 if data[0][0]=='Kerros' and n==5:weights=[.13,.20,.20,.24,.23]
 if data[0][0]=='Materiaali':weights=[.20,.36,.22,.22]
 if data[0][0]=='Tarkistettava':weights=[.55,.08,.08,.29]
 if data[0][0]=='Tarkistettava tieto':weights=[.18,.22,.38,.22]
 if data[0][0]=='Elementti tai toimitusrivi':weights=[.18,.10,.27,.23,.22]
 if data[0][0]=='Aiempi nimike' and n==3:weights=[.29,.53,.18]
 for col,w in zip(t.columns,weights):col.width=Cm(17.4*w)
 for i,values in enumerate(data):
  row=t.add_row()
  pr=row._tr.get_or_add_trPr();avoid=OxmlElement('w:cantSplit');pr.append(avoid)
  if i==0:
   repeat=OxmlElement('w:tblHeader');pr.append(repeat)
  for j,(cell,value) in enumerate(zip(row.cells,values)):
   cell.width=Cm(17.4*weights[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tc.append(b);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');m.append(e)
   tc.append(m)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F8E0C7');tc.append(sh)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
   p.paragraph_format.keep_with_next=(len(data)<=8 and i<len(data)-1)
   r=p.add_run(value);r.font.name='Work Sans';r.font.size=Pt(9);r.bold=i==0
 d.add_paragraph().paragraph_format.space_after=Pt(1)
files=sorted(root.glob('[0-9]*.md'))
for index,path in enumerate(files):
 if index:pass
 lines=path.read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i];i+=1
  if not line.strip():continue
  if line.startswith('|'):
   rows=[line]
   while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
   table(rows);continue
  if line.startswith('# '):
   title=clean(line[2:]);p=d.add_paragraph(title,'Title' if index==0 else 'Heading 1');p.paragraph_format.page_break_before=bool(index);continue
  if line.startswith('## '):
   if path.name.startswith('07-') and line[3:] in ['Jorma Salomäelle','Tapani Katajistolle']:d.add_page_break()
   d.add_paragraph(clean(line[3:]),'Heading 2');continue
  if line.startswith('### '):d.add_paragraph(clean(line[4:]),'Heading 3');continue
  p=d.add_paragraph(clean(line))
  if line.startswith(('ASIAKASTEKSTI','LÄHDETIETO','TARKISTETTAVA','EHDOTUS')):
   p.paragraph_format.space_before=Pt(4)
  if path.name.startswith('02-'):p.paragraph_format.space_after=Pt(5)
for para in d.paragraphs + list(sec.header.paragraphs) + list(sec.footer.paragraphs):
 for run in para.runs:run.font.name='Work Sans'
d.core_properties.title='Hietakulman dokumenttien ja liitteiden tarkistuspaketti'
d.core_properties.subject='Luonnos 0.3 Villen Jorman ja Tapanin tarkistettavaksi'
d.core_properties.author='Hietakulma'
out=root/'Hietakulma-tarkistuspaketti-LUONNOS-0.3.docx';d.save(out);print(out)
